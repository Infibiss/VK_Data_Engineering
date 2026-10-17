"""Prepare the requested concise revision from the recorded HW2 evidence."""
from pathlib import Path
import hashlib
import json
import shutil
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'output/docx/HW2_report.docx'
BACKUP = ROOT / 'tmp/HW2_report_before_gost.docx'
if OUTPUT.exists() and not BACKUP.exists():
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUTPUT, BACKUP)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin, sec.right_margin = Mm(30), Mm(15)
sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)
sec.footer_distance = Mm(10)
sec.different_first_page_header_footer = True

def font(style, family='Times New Roman', size=14, bold=False):
    style.font.name = family
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr()
    families = rf.find(qn('w:rFonts'))
    if families is None:
        families = OxmlElement('w:rFonts')
        rf.append(families)
    for attr in tuple(families.attrib):
        if 'theme' in attr.lower():
            del families.attrib[attr]
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        families.set(qn('w:' + attr), family)
    complex_size = rf.find(qn('w:szCs'))
    if complex_size is not None:
        complex_size.set(qn('w:val'), str(size * 2))

normal = doc.styles['Normal']
font(normal)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.first_line_indent = Cm(1.25)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.widow_control = True

for name in ('Heading 1', 'Heading 2'):
    st = doc.styles[name]
    font(st, bold=True)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.first_line_indent = Cm(1.25)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_before = Pt(0 if name == 'Heading 1' else 8)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.keep_together = True
    st.paragraph_format.page_break_before = name == 'Heading 1'

font(doc.styles['Title'], size=14, bold=True)
for border in list(doc.styles['Title'].element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
for spacing in list(doc.styles['Title'].element.get_or_add_rPr().iter(qn('w:spacing'))):
    spacing.getparent().remove(spacing)
doc.styles['Title'].paragraph_format.first_line_indent = Cm(0)
doc.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Title'].paragraph_format.line_spacing = 1.5
doc.styles['Title'].paragraph_format.space_after = Pt(0)

for name, family, size in (
    ('Code', 'Courier New', 12),
    ('Figure Caption', 'Times New Roman', 12),
    ('Table Caption', 'Times New Roman', 12),
    ('Table Text', 'Times New Roman', 12),
    ('Reference', 'Times New Roman', 14),
):
    st = doc.styles.add_style(name, 1)
    st.base_style = normal
    font(st, family=family, size=size)
    st.paragraph_format.first_line_indent = Cm(0)
    st.paragraph_format.line_spacing = 1.5 if name == 'Reference' else 1.0
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.space_after = Pt(6 if name != 'Table Text' else 0)
    st.paragraph_format.keep_together = True

doc.styles['Figure Caption'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Table Caption'].paragraph_format.keep_with_next = True
doc.styles['Table Caption'].paragraph_format.space_before = Pt(7)

# Page numbering includes the title page; its first-page footer remains empty.
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent = Cm(0)
footer.paragraph_format.line_spacing = 1.0
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
fr = OxmlElement('w:r')
ft = OxmlElement('w:t')
ft.text = '2'
fr.append(ft)
field.append(fr)
footer._p.append(field)

doc.core_properties.title = 'Отчёт по заданию HW2 RustFS и Erasure Coding'
doc.core_properties.author = 'Муртазалиев М. А.'
doc.core_properties.subject = 'Группа J3310, ИСУ 466797'
doc.core_properties.keywords = 'HW2, RustFS, EC 2+2'

def paragraph(text, style=None):
    return doc.add_paragraph(text, style=style)

def heading(text, level=1, new_page=False):
    p = doc.add_paragraph(text, style=f'Heading {level}')
    if new_page:
        p.paragraph_format.page_break_before = True
    return p

def code(text):
    return paragraph(text, 'Code')

def math_run(text):
    r = OxmlElement('m:r')
    rp = OxmlElement('m:rPr')
    sty = OxmlElement('m:sty')
    sty.set(qn('m:val'), 'p')
    rp.append(sty)
    r.append(rp)
    wrp = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '28')
    wrp.append(sz)
    r.append(wrp)
    t = OxmlElement('m:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    return r

def fraction(top, bottom):
    f = OxmlElement('m:f')
    for tag, value in (('m:num', top), ('m:den', bottom)):
        e = OxmlElement(tag)
        e.append(math_run(value))
        f.append(e)
    return f

def subscript(base, sub):
    e = OxmlElement('m:sSub')
    for tag, value in (('m:e', base), ('m:sub', sub)):
        item = OxmlElement(tag)
        item.append(math_run(value))
        e.append(item)
    return e

def equation(*parts):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    om = OxmlElement('m:oMath')
    for part in parts:
        om.append(math_run(part) if isinstance(part, str) else part)
    p._p.append(om)

def table(caption, rows, widths):
    paragraph(caption, 'Table Caption')
    t = doc.add_table(rows=0, cols=len(widths))
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Mm(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        for key, value in (('val', 'single'), ('sz', '6'), ('color', 'D9D9D9')):
            el.set(qn('w:' + key), value)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i, values in enumerate(rows):
        row = t.add_row()
        props = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        props.append(no_split)
        if i == 0:
            props.append(OxmlElement('w:tblHeader'))
        for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, val in (('top', '75'), ('bottom', '75'), ('left', '85'), ('right', '85')):
                el = OxmlElement('w:' + side)
                el.set(qn('w:w'), val)
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            tcpr.append(margins)
            if i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'F2F2F2')
                tcpr.append(shade)
            p = cell.paragraphs[0]
            p.style = doc.styles['Table Text']
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(str(value))
    return t

photos = []
def photo(file, caption, width=165):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ROOT / 'evidence' / file), width=Mm(width))
    paragraph(caption, 'Figure Caption')
    photos.append(file)

def link(p, label, url):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    rel = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), rel)
    r = OxmlElement('w:r')
    rp = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '000000')
    rp.append(color)
    r.append(rp)
    t = OxmlElement('w:t')
    t.text = label
    r.append(t)
    h.append(r)
    p._p.append(h)

def fmt(n):
    return f'{n:,}'.replace(',', ' ')

# Title: only the supplied personal data and the subject of the assignment.
p = paragraph('ОТЧЁТ ПО ЗАДАНИЮ HW2', 'Title')
p.paragraph_format.space_before = Pt(150)
paragraph('RustFS и Erasure Coding', 'Title')
p = paragraph('Выполнил: Муртазалиев М. А.')
p.paragraph_format.space_before = Pt(100)
for p in (p, paragraph('Группа J3310'), paragraph('ИСУ 466797')):
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p = paragraph('2026')
p.paragraph_format.first_line_indent = Cm(0)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(130)

heading('1 Развёртывание и загрузка данных')
heading('1.1 Развёртывание RustFS', level=2)
paragraph('RustFS 1.0.1 запущен на macOS arm64: один узел, четыре папки-диска volumes/disk1–disk4, схема EC 2+2. S3 API: 127.0.0.1:9000; консоль: 127.0.0.1:9001; бакет hw2. Состояние кластера показано на рисунке 1.')
paragraph('Папки размещены на одном SSD. Для этой модели включён RUSTFS_UNSAFE_BYPASS_DISK_CHECK; испытания проверяют отказ путей к папкам, а не физических дисков.')
code('RUSTFS_STORAGE_CLASS_STANDARD=EC:2\n'
     'RUSTFS_ERASURE_SET_DRIVE_COUNT=4\n'
     'RUSTFS_UNSAFE_BYPASS_DISK_CHECK=true\n'
     '$ zsh scripts/start.sh\n'
     '$ ./bin/rustfs --version\n'
     'rustfs 1.0.1')
photo('01-cluster.jpg', 'Рисунок 1 – Кластер RustFS с четырьмя доступными папками')

heading('1.2 Загрузка данных', level=2, new_page=True)
paragraph('Загружен датасет ИТ-вакансий it_vacancy_data [1]: 351 623 записи. Три JSON-файла сформированы из первых трёх записей CSV. Объекты и их размеры приведены в таблице 1; загрузка подтверждена рисунком 2.')
manifest = json.loads((ROOT / 'evidence/manifest.json').read_text())
table('Таблица 1 – Загруженные объекты',
      [['Объект', 'Размер, байт', 'Загрузка PUT, HTTP']] +
      [[x['file'], fmt(x['bytes']), '200'] for x in manifest['files']],
      [78, 42, 45])
code('$ python3 scripts/load_data.py')
photo('02-objects.jpg', 'Рисунок 2 – Датасет CSV и три JSON-файла в бакете hw2')

heading('2 Размеры папок и механизм Erasure Coding')
paragraph('Размеры измерены командой du -sk до испытаний отказов (таблица 2). 1 КиБ = 1024 байта; du учитывает выделенные блоки файловой системы.')
code('$ du -sk volumes/disk*\n'
     '42856  volumes/disk1\n42856  volumes/disk2\n'
     '42856  volumes/disk3\n42856  volumes/disk4')
table('Таблица 2 – Размеры папок и фрагментов CSV',
      [['Папка', 'du, байт', 'Сумма файлов, байт', 'part.1 CSV, байт']] +
      [[f'disk{i}', '43 884 544', '43 801 241', '43 780 467'] for i in range(1, 5)],
      [20, 45, 50, 50])
paragraph('В xl.meta CSV: EcM=2, EcN=2, EcBSize=1 048 576, EcDist=[2,3,4,1]. disk1 хранит D2, disk2 — P1, disk3 — P2, disk4 — D1. Reed–Solomon делит блок на фрагменты данных D1 и D2; чётность P1 и P2 — их кодовые комбинации. Для восстановления достаточно любых двух фрагментов [2].')
paragraph('CSV содержит 83 полных блока по 1 МиБ и остаток 523 750 байт. В полном блоке D1, D2, P1 и P2 имеют по 524 288 байт; в последнем — по 261 875 байт.')
equation('83 × 1 048 576 + 523 750 = 87 555 558')
equation('83 × 524 288 + 261 875 = 43 777 779')
paragraph('В part.1 записаны 84 фрагмента блоков и их 32-байтовые контрольные суммы: 84 × 32 = 2 688 байт. Поэтому part.1 на каждой папке занимает 43 777 779 + 2 688 = 43 780 467 байт, что совпадает с замером.')
paragraph('Фрагменты JSON хранятся внутри xl.meta. Сумма файлов четырёх папок — 175 204 964 байта: примерно удвоенный объём объектов с учётом метаданных и контрольных сумм.')

heading('3 Испытания отказов')
paragraph('Каждый сценарий начат с четырёх доступных папок. Скрипт переносит выбранные папки в offline и заменяет их исходные пути обычными файлами, исключая автоматическое создание новых папок. Ниже приведены фрагменты вывода проверки чтения.')

for n in (1, 2, 3):
    heading({1: '3.1 Потеря одной папки', 2: '3.2 Потеря половины папок',
             3: '3.3 Потеря половины плюс одной папки'}[n], level=2, new_page=n != 1)
    source = (ROOT / 'evidence' / {1: '04-loss-1.txt', 2: '05-loss-2.txt', 3: '06-loss-3.txt'}[n]).read_text()
    lines = source.splitlines()
    selected = [f'$ python3 scripts/failure.py disable {n}',
                f'$ python3 scripts/failure.py check {n}',
                f'Unavailable drive paths: {n}/4']
    for i, line in enumerate(lines):
        if line.startswith('$ python3 scripts/s3.py GET '):
            selected.extend([line, lines[i + 1].replace('; SHA256 MATCH', '')])
    code('\n'.join(selected))
    if n == 1:
        paragraph('Недоступна disk1: остались 3 фрагмента из 4. Все объекты прочитаны (HTTP 200); консоль показывает 3 Online и 1 Offline (рисунок 3).')
    elif n == 2:
        paragraph('Недоступны disk1 и disk2: остались D1 и P2. Двух фрагментов достаточно; все объекты прочитаны (HTTP 200). В консоли 2 Online и 2 Offline (рисунок 4).')
    else:
        paragraph('Недоступны disk1–disk3: остался только D1. Все чтения вернули HTTP 503, SlowDownRead; 162 байта — сообщение об ошибке. В консоли 1 Online и 3 Offline (рисунок 5).')
    photo({1: '04-loss-1.jpg', 2: '05-loss-2.jpg', 3: '06-loss-3.jpg'}[n],
          {1: 'Рисунок 3 – Потеря одной папки из четырёх (25 %)',
           2: 'Рисунок 4 – Потеря двух папок из четырёх (50 %)',
           3: 'Рисунок 5 – Потеря трёх папок из четырёх (75 %)'}[n], width=160)

heading('4 Расчёт стоимости хранения')
paragraph('Для схемы EC 2+2 K=2 фрагмента данных, M=2 фрагмента чётности. Коэффициент избыточности равен отношению физического объёма к исходному. Для трёх полных копий коэффициент равен 3.')
equation(subscript('α', 'EC'), ' = ', fraction('K + M', 'K'),
         ' = ', fraction('4', '2'), ' = 2;    ', subscript('α', '3'), ' = 3')
paragraph('Пусть V — исходный объём в ГБ, p — цена хранения 1 ГБ за расчётный период. Тогда физический объём и стоимость определяются формулами:')
equation(subscript('V', 'физ'), ' = αV;    C = αVp')
paragraph('Тариф в задании не указан; суммы выражены через p. Использованы десятичные единицы: 1 ТБ = 1000 ГБ. Результаты приведены в таблице 3.')
table('Таблица 3 – Объём и стоимость хранения',
      [['Исходный объём', 'EC 2+2', 'Три копии'],
       ['25 ГБ', '50 ГБ\n50p', '75 ГБ\n75p'],
       ['50 ТБ = 50 000 ГБ', '100 ТБ\n100 000p', '150 ТБ\n150 000p']],
      [61, 52, 52])
equation('Экономия = ', fraction('3 − 2', '3'), ' × 100 % = 33,33 %')
paragraph('EC экономит 25 ГБ для исходных 25 ГБ и 50 ТБ для исходных 50 ТБ. Расчёт выполнен без служебных метаданных.')

heading('5 Обоснование выбора хранилища')
paragraph('Для хранения датасета CSV и JSON-файлов выбран RustFS со схемой EC 2+2 и доступом через S3 API.')
paragraph('Надёжность: чтение сохраняется при потере любых двух фрагментов из четырёх. Испытания подтвердили доступность при отказе одной и двух папок и недоступность при отказе трёх.')
paragraph('Стоимость: EC требует двух исходных объёмов вместо трёх при репликации; экономия составляет 33,33 %.')
paragraph('Производительность: EC требует вычисления чётности при записи и реконструкции данных при потере фрагментов. Три копии позволяют читать объект с одной реплики и восстанавливать его копированием, но требуют большего объёма хранения. Выбор EC обоснован экономией ёмкости при сохранении чтения после двух отказов; скорость в этом опыте не измерялась.')

p = heading('СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Cm(0)
p = paragraph('1. Ternikov A. it_vacancy_data. Figshare. 2022. URL: ', 'Reference')
link(p, 'figshare.com/articles/dataset/it_vacancy_data/19005092',
     'https://figshare.com/articles/dataset/it_vacancy_data/19005092')
p.add_run(' (дата обращения: 05.10.2026).')
p = paragraph('2. RustFS. Erasure Coding Architecture. Версия 1.0.1. URL: ', 'Reference')
link(p, 'github.com/rustfs/rustfs/blob/1.0.1/docs/architecture/erasure-coding.md',
     'https://github.com/rustfs/rustfs/blob/1.0.1/docs/architecture/erasure-coding.md')
p.add_run(' (дата обращения: 05.10.2026).')

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)

# Record evidence provenance separately, outside the submitted report.
with ZipFile(OUTPUT) as z:
    media = {hashlib.sha256(z.read(name)).hexdigest(): name
             for name in z.namelist() if name.startswith('word/media/')}
matched = []
for name in photos:
    digest = hashlib.sha256((ROOT / 'evidence' / name).read_bytes()).hexdigest()
    assert digest in media, name
    matched.append({'evidence': name, 'docx_media': media[digest], 'unaltered': True})
(ROOT / 'evidence/docx-gost-screenshot-verification.json').write_text(
    json.dumps(matched, ensure_ascii=False, indent=2))
print(OUTPUT)
print(f'Original screenshot files embedded unchanged: {len(matched)}')
